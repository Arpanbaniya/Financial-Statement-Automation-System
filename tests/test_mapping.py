"""Check conservative aliases, review decisions, and source preservation."""

from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook

from ingestion import ingest_document
from normalization import SourceLocation, correct_mapping, map_label, map_table_rows

FIXTURES = Path(__file__).parent / "fixtures"
REVIEWER = "11111111-1111-4111-8111-111111111111"
LOCATION = SourceLocation(sheet="Income", cell="A4", row=4)


@pytest.mark.parametrize("label", ["Net sales", "Sales", "Turnover"])
def test_exact_income_aliases_have_source_evidence(label: str) -> None:
    record = map_label(label, "income_statement", source_location=LOCATION)
    assert record.source_label == label
    assert record.canonical_field == "revenue"
    assert record.method == "exact_alias"
    assert record.confidence == 0.97
    assert record.status == "suggested"
    assert record.source_location == LOCATION
    assert record.evidence[0].matched_alias == label
    assert record.reviewer_id is None
    assert record.reviewed_at is None


@pytest.mark.parametrize(
    ("label", "expected"),
    [
        ("NET-SALES:", "revenue"),
        ("cost.of sales", "cost_of_revenue"),
        ("  STOCKHOLDERS’ EQUITY  ", "shareholders_equity"),
        ("Selling, General & Administrative", "selling_general_administrative"),
    ],
)
def test_case_punctuation_and_spacing_use_normalized_aliases(
    label: str, expected: str
) -> None:
    statement_type = (
        "balance_sheet" if expected == "shareholders_equity" else "income_statement"
    )
    record = map_label(label, statement_type, source_location=LOCATION)
    assert record.canonical_field == expected
    assert record.method == "normalized_alias"
    assert record.confidence == 0.90
    assert record.evidence[0].compared_text


def test_generic_other_and_unmatched_labels_are_sent_to_review() -> None:
    other = map_label("Other", "balance_sheet", source_location=LOCATION)
    assert other.canonical_field is None
    assert other.status == "needs_review"
    assert other.review_reason == "ambiguous_label"
    assert {
        "other_current_assets",
        "other_noncurrent_assets",
    } <= set(other.candidate_fields)

    assets = map_label("Other assets", "balance_sheet", source_location=LOCATION)
    assert assets.canonical_field is None
    assert assets.review_reason == "ambiguous_label"
    assert "other_current_assets" in assets.candidate_fields
    assert "other_noncurrent_assets" in assets.candidate_fields

    specific = map_label(
        "Other current assets", "balance_sheet", source_location=LOCATION
    )
    assert specific.canonical_field == "other_current_assets"


@pytest.mark.parametrize(
    "label",
    [
        "Revenue growth",
        "Reveneu",
        "Sales from discontinued operations",
        "Interest income tax",
        "Operating costs",
        "Finance income",
    ],
)
def test_partial_phrases_and_typos_are_not_silently_mapped(label: str) -> None:
    record = map_label(label, "income_statement", source_location=LOCATION)
    assert record.canonical_field is None
    assert record.status == "needs_review"
    assert record.review_reason == "unmatched_label"
    assert record.confidence == 0


def test_very_long_labels_are_preserved_but_not_matched() -> None:
    label = "Revenue" + " " * 501
    record = map_label(label, "income_statement", source_location=LOCATION)
    assert record.source_label == label
    assert record.canonical_field is None
    assert record.review_reason == "unmatched_label"


def test_broad_payables_label_is_not_reduced_to_trade_payables() -> None:
    record = map_label(
        "Trade and other payables", "balance_sheet", source_location=LOCATION
    )
    assert record.canonical_field is None
    assert record.review_reason == "unmatched_label"


def test_statement_context_and_unknown_type_prevent_cross_mapping() -> None:
    income = map_label("Net income", "income_statement", source_location=LOCATION)
    cash = map_label("Net income", "cash_flow_statement", source_location=LOCATION)
    assert income.canonical_field == cash.canonical_field == "net_income"
    assert income.statement_type != cash.statement_type

    unknown = map_label("Net income", "unknown", source_location=LOCATION)
    assert unknown.canonical_field is None
    assert unknown.review_reason == "unknown_statement"

    empty = map_label("  ", "income_statement", source_location=LOCATION)
    assert empty.review_reason == "empty_label"


def test_colliding_aliases_are_not_silently_selected(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import normalization.mapping as mapping

    index = {key: dict(value) for key, value in mapping._EXACT_INDEX.items()}
    index["income_statement"]["Sales"] = [
        ("revenue", "Sales"),
        ("net_income", "Sales"),
    ]
    monkeypatch.setattr(mapping, "_EXACT_INDEX", index)
    record = map_label("Sales", "income_statement", source_location=LOCATION)
    assert record.canonical_field is None
    assert record.review_reason == "ambiguous_label"
    assert record.candidate_fields == ("net_income", "revenue")
    assert {item.matched_alias for item in record.evidence} == {"Sales"}


def test_table_mapping_keeps_original_labels_and_coordinates() -> None:
    source = (FIXTURES / "sample_statement.csv").read_bytes()
    extracted = ingest_document(source, filename="statement.csv")
    records = map_table_rows(extracted.tables[0], "income_statement")
    assert len(records) == 3
    assert records[0].review_reason == "unmatched_label"
    assert records[1].canonical_field == "revenue"
    assert records[1].source_location == SourceLocation(
        cell="A2", table="CSV records", row=2, line=2
    )
    assert records[2].source_label == "Cost of\nsales"
    assert records[2].canonical_field == "cost_of_revenue"
    assert records[2].method == "normalized_alias"
    assert extracted.tables[0].rows[2].cells[0].value == "Cost of\nsales"
    assert source == (FIXTURES / "sample_statement.csv").read_bytes()


def test_workbook_table_mapping_keeps_sheet_and_cell() -> None:
    from io import BytesIO

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Income"
    sheet["B4"] = "Net sales"
    sheet["C4"] = 100
    output = BytesIO()
    workbook.save(output)
    extracted = ingest_document(output.getvalue(), filename="income.xlsx")
    record = map_table_rows(extracted.tables[0], "income_statement")[0]
    assert record.canonical_field == "revenue"
    assert record.source_location == SourceLocation(
        sheet="Income", cell="B4", table="Income region 1", row=4
    )


def test_manual_correction_records_reviewer_time_without_mutating_source() -> None:
    original = map_label("Other assets", "balance_sheet", source_location=LOCATION)
    local_time = datetime(
        2026, 9, 26, 12, 30, tzinfo=timezone(timedelta(hours=5, minutes=45))
    )
    corrected = correct_mapping(
        original,
        statement_type="balance_sheet",
        canonical_field="other_current_assets",
        reviewer_id=REVIEWER,
        reviewed_at=local_time,
    )
    assert corrected is not original
    assert original.canonical_field is None
    assert original.status == "needs_review"
    assert corrected.canonical_field == "other_current_assets"
    assert corrected.method == "manual"
    assert corrected.status == "accepted"
    assert corrected.confidence == 1
    assert corrected.reviewer_id == REVIEWER
    assert corrected.reviewed_at == local_time.astimezone(UTC)
    assert corrected.source_label == original.source_label
    assert corrected.source_location == original.source_location
    assert corrected.evidence[-1].method == "manual"

    with pytest.raises(KeyError):
        correct_mapping(
            original,
            statement_type="balance_sheet",
            canonical_field="revenue",
            reviewer_id=REVIEWER,
        )
    with pytest.raises(ValueError, match="UUID"):
        correct_mapping(
            original,
            statement_type="balance_sheet",
            canonical_field="total_assets",
            reviewer_id="person",
        )
    with pytest.raises(ValueError, match="time zone"):
        correct_mapping(
            original,
            statement_type="balance_sheet",
            canonical_field="total_assets",
            reviewer_id=REVIEWER,
            reviewed_at=datetime(2026, 9, 26),
        )


def test_location_rejects_nonpositive_references() -> None:
    with pytest.raises(ValueError):
        SourceLocation(page=0)
    with pytest.raises(ValueError):
        SourceLocation(row=-1)
    with pytest.raises(ValueError):
        SourceLocation(line=0)
