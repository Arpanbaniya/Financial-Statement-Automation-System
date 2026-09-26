"""Open the generated report to check sheets, values, and provenance."""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from finance import calculate_cash_flow, calculate_ratios
from normalization import SourceLocation, normalize_period
from reports import SHEETS, build_excel_report, report_filename
from validation import SourceRef, StatementSnapshot, ValidationValue

NOW = datetime(2026, 9, 26, tzinfo=UTC)


def statement(kind: str, name: str, amounts: dict[str, str]) -> StatementSnapshot:
    return StatementSnapshot(
        id=name,
        company_id="company-1",
        document_id="document-1",
        statement_type=kind,  # type: ignore[arg-type]
        period=normalize_period(
            kind,  # type: ignore[arg-type]
            start="2025-01-01" if kind != "balance_sheet" else None,
            end="2025-12-31",
        ),
        values=tuple(
            ValidationValue(
                field=field,
                normalized_value=Decimal(amount),
                original_value=amount,
                source_ref=SourceRef(
                    name, field, SourceLocation(sheet="FY25", cell=f"B{index}")
                ),
                currency="USD",
                unit_scale="ones",
                status="accepted",
            )
            for index, (field, amount) in enumerate(amounts.items(), 2)
        ),
        currency="USD",
        unit_scale="ones",
    )


def test_workbook_has_all_sheets_values_and_source_links() -> None:
    income = statement(
        "income_statement", "income", {"revenue": "1000", "net_income": "80"}
    )
    balance = statement(
        "balance_sheet",
        "balance",
        {"total_current_assets": "600", "total_current_liabilities": "300"},
    )
    cash = statement(
        "cash_flow_statement",
        "cash",
        {"operating_cash_flow": "120", "capital_expenditure": "-30"},
    )
    data = build_excel_report(
        company_name="Example Ltd",
        statements=(income, balance, cash),
        ratios=calculate_ratios(
            income_statement=income,
            ending_balance_sheet=balance,
            calculated_at=NOW,
        ),
        cash_flow=calculate_cash_flow(
            cash_flow_statement=cash,
            income_statement=income,
            calculated_at=NOW,
        ),
        generated_at=NOW,
    )
    book = load_workbook(BytesIO(data), data_only=False)
    assert tuple(book.sheetnames) == SHEETS
    assert book["Summary"]["A1"].value == "Financial analysis"
    assert book["Summary"].freeze_panes == "B5"
    assert book["Cash Flow"]["C5"].value == 120
    assert book["Source Data"]["D5"].value == "1000"
    assert book["Source Data"]["E5"].value == 1000
    assert "FY25!B2" in book["Source Data"]["H5"].value
    assert any(
        row[0].value == "Free Cash Flow" and row[1].value == 90
        for row in book["Summary"].iter_rows(min_row=5)
    )
    assert any(
        row[0].value == "free_cash_flow"
        for row in book["Methodology"].iter_rows(min_row=5)
    )
    assert not book._external_links


def test_untrusted_company_name_cannot_become_excel_formula() -> None:
    income = statement("income_statement", "income", {"revenue": "1"})
    data = build_excel_report(
        company_name='=HYPERLINK("https://example.test")',
        statements=(income,),
        generated_at=NOW,
    )
    book = load_workbook(BytesIO(data))
    assert book["Summary"]["A2"].data_type == "s"
    assert (
        report_filename("../Acme Inc", "2025-12-31")
        == "Financial_Analysis_Acme_Inc_2025-12-31.xlsx"
    )
    with pytest.raises(ValueError, match="timezone"):
        build_excel_report(
            company_name="Example",
            statements=(income,),
            generated_at=datetime(2026, 9, 26),
        )
