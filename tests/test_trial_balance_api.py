"""Authorization, generation writes, recomputation and download integration."""

import hashlib
import json
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from api.index import app
from tests.test_trial_balance import COMPANY, DOCUMENT, SOURCE, cash, request

USER = "33333333-3333-4333-8333-333333333333"
OTHER = "44444444-4444-4444-8444-444444444444"


@pytest.fixture
def service(monkeypatch):
    monkeypatch.setenv("NEXT_PUBLIC_SUPABASE_URL", "https://example.supabase.co")
    monkeypatch.setenv("NEXT_PUBLIC_SUPABASE_PUBLISHABLE_KEY", "sb_publishable_test")
    monkeypatch.setenv("SUPABASE_SECRET_KEY", "sb_secret_test")
    monkeypatch.setenv("AI_PROVIDER", "none")
    state = {"writes": [], "source": SOURCE, "save_status": 200}
    original = httpx.AsyncClient

    def handle(req):
        path = req.url.path
        if path == "/auth/v1/user":
            token = req.headers.get("Authorization")
            return httpx.Response(
                200, json={"id": USER if token == "Bearer owner" else OTHER}
            )
        if path == "/rest/v1/rpc/save_trial_balance":
            state["writes"].append(json.loads(req.content))
            return httpx.Response(state["save_status"], json=None)
        if path.startswith("/rest/v1/"):
            assert req.url.params["user_id"] in {f"eq.{USER}", f"eq.{OTHER}"}
            if req.url.params["user_id"] == f"eq.{OTHER}":
                return httpx.Response(200, json=[])
            if path.endswith("/companies"):
                return httpx.Response(
                    200, json=[{"id": COMPANY, "name": "Synthetic Co"}]
                )
            if path.endswith("/trial_balance_runs"):
                return httpx.Response(200, json=[])
            if path.endswith("/documents"):
                return httpx.Response(
                    200,
                    json=[
                        {
                            "id": DOCUMENT,
                            "storage_path": f"{USER}/{DOCUMENT}/tb.csv",
                            "original_filename": "tb.csv",
                            "mime_type": "text/csv",
                            "file_size": len(state["source"]),
                            "sha256": hashlib.sha256(state["source"]).hexdigest(),
                            "status": "uploaded",
                        }
                    ],
                )
        if path.startswith("/storage/v1/object/"):
            return httpx.Response(200, content=state["source"])
        raise AssertionError(path)

    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(handle), **kwargs),
    )
    with TestClient(app) as client:
        yield client, state


def test_authentication_and_ownership(service):
    client, state = service
    assert client.get(f"/api/trial-balance/{DOCUMENT}").status_code == 401
    assert (
        client.get(
            f"/api/trial-balance/{DOCUMENT}", headers={"Authorization": "Bearer other"}
        ).status_code
        == 404
    )
    assert (
        client.post(
            f"/api/trial-balance/{DOCUMENT}/generate",
            json=request().model_dump(mode="json"),
            headers={"Authorization": "Bearer other"},
        ).status_code
        == 404
    )
    assert state["writes"] == []


def test_generation_uses_single_atomic_write_and_stable_ids(service):
    client, state = service
    payload = request(opening_complete=True, cash=cash()).model_dump(mode="json")
    for _ in range(2):
        response = client.post(
            f"/api/trial-balance/{DOCUMENT}/generate",
            json=payload,
            headers={"Authorization": "Bearer owner"},
        )
        assert response.status_code == 200, response.text
        assert response.json()["saved"]
        assert len(response.json()["statements"]) == 3
    assert len(state["writes"]) == 2
    assert state["writes"][0]["p_statements"] == state["writes"][1]["p_statements"]
    assert len(state["writes"][0]["p_metrics"]) > 5
    assert (
        sum(row["status"] == "accepted" for row in state["writes"][0]["p_statements"])
        == 3
    )


def test_imbalance_blocks_write_and_save_failure_is_visible(service):
    client, state = service
    state["source"] = SOURCE.replace(b"Cash,800", b"Cash,900")
    payload = request().model_dump(mode="json")
    response = client.post(
        f"/api/trial-balance/{DOCUMENT}/generate",
        json=payload,
        headers={"Authorization": "Bearer owner"},
    )
    assert response.status_code == 422
    assert not state["writes"]
    state["source"] = SOURCE
    state["save_status"] = 500
    assert (
        client.post(
            f"/api/trial-balance/{DOCUMENT}/generate",
            json=payload,
            headers={"Authorization": "Bearer owner"},
        ).status_code
        == 502
    )


def test_excel_and_deterministic_explanation_do_not_write(service):
    client, state = service
    payload = request().model_dump(mode="json")
    headers = {"Authorization": "Bearer owner"}
    response = client.post(
        f"/api/trial-balance/{DOCUMENT}/excel", json=payload, headers=headers
    )
    assert response.status_code == 200
    assert "Trial Balance" in load_workbook(BytesIO(response.content)).sheetnames
    explanation = client.post(
        f"/api/trial-balance/{DOCUMENT}/explain", json=payload, headers=headers
    )
    assert explanation.status_code == 200
    assert explanation.json()["explanation"]["fallback_used"]
    assert state["writes"] == []
