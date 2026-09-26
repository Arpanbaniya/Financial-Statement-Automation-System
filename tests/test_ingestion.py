"""Check that each adapter keeps source values and reports recoverable problems."""

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from pypdf import PdfWriter

from ingestion import IngestionError, ingest_document

FIXTURES = Path(__file__).parent / "fixtures"


def pdf_with_text(text: str) -> bytes:
    """Build a tiny text PDF fixture without needing a PDF editor in production."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length "
        + str(len(stream)).encode()
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_start = len(content)
    content.extend(f"xref\n0 {len(offsets)}\n".encode())
    content.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode())
    trailer = (
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{xref_start}\n%%EOF"
    )
    content.extend(trailer.encode())
    return bytes(content)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Income statement"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 100
    sheet["B2"] = "=B1-40"
    workbook.create_sheet("Notes")["C3"] = "Source note"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_csv_preserves_text_delimiter_records_and_lines() -> None:
    source = (FIXTURES / "sample_statement.csv").read_bytes()
    original = bytes(source)
    result = ingest_document(source, filename="statement.csv", mime_type="text/csv")

    assert source == original
    assert result.file_type == "csv"
    assert result.raw_text == source.decode("utf-8")
    assert result.tables[0].delimiter == ";"
    assert result.tables[0].rows[2].cells[0].value == "Cost of\nsales"
    assert result.tables[0].rows[2].source_line == 4
    assert result.tables[0].rows[1].cells[1].coordinate == "B2"
    assert result.extraction_method == "python-csv"
    assert result.extraction_version


def test_csv_reports_irregular_rows_and_bad_encoding() -> None:
    result = ingest_document(b"A,B\n1\n", filename="odd.csv")
    assert any(warning.code == "COLUMN_COUNT_MISMATCH" for warning in result.warnings)
    with pytest.raises(IngestionError, match="UTF-8") as error:
        ingest_document(b"\xff\xfe", filename="old.csv")
    assert error.value.code == "UNSUPPORTED_ENCODING"


def test_xlsx_preserves_sheet_cell_coordinates_types_and_formulas() -> None:
    source = workbook_bytes()
    original = bytes(source)
    result = ingest_document(source, filename="statement.xlsx")

    assert source == original
    assert result.file_type == "xlsx"
    assert [sheet.name for sheet in result.sheets] == ["Income statement", "Notes"]
    first_row = result.sheets[0].rows[0]
    assert first_row.cells[0].coordinate == "A1"
    assert first_row.cells[1].value == 100
    assert result.sheets[0].rows[1].cells[0].value == "=B1-40"
    assert result.sheets[0].rows[1].cells[0].data_type == "f"
    assert result.sheets[1].rows[0].cells[0].coordinate == "C3"
    assert result.extraction_method == "openpyxl-read-only"
    assert result.extraction_version


def test_pdf_keeps_page_text_and_warns_for_image_only_page() -> None:
    text_result = ingest_document(pdf_with_text("Revenue 100"), filename="text.pdf")
    assert text_result.pages[0].number == 1
    assert "Revenue 100" in text_result.pages[0].text
    assert text_result.raw_text == text_result.pages[0].text
    assert text_result.extraction_method == "pypdf-text"
    assert text_result.extraction_version

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    output = BytesIO()
    writer.write(output)
    blank_result = ingest_document(output.getvalue(), filename="scan.pdf")
    assert blank_result.pages[0].text == ""
    assert blank_result.warnings[0].code == "NO_EXTRACTABLE_TEXT"


def test_encrypted_pdf_has_a_clear_failure_code() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("example-password")
    output = BytesIO()
    writer.write(output)
    with pytest.raises(IngestionError) as error:
        ingest_document(output.getvalue(), filename="locked.pdf")
    assert error.value.code == "ENCRYPTED_PDF"


def test_xlsx_archive_expansion_is_bounded(monkeypatch: pytest.MonkeyPatch) -> None:
    from ingestion import reader

    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("large.xml", b"x" * 101)
    monkeypatch.setattr(reader, "MAX_UNCOMPRESSED_XLSX_BYTES", 100)
    with pytest.raises(IngestionError) as error:
        ingest_document(output.getvalue(), filename="large.xlsx")
    assert error.value.code == "DOCUMENT_TOO_COMPLEX"


@pytest.mark.parametrize(
    ("source", "filename", "code"),
    [
        (b"", "file.pdf", "EMPTY_SOURCE"),
        (b"hello", "file.exe", "UNSUPPORTED_FILE_TYPE"),
        (b"hello", "file.pdf", "INVALID_PDF"),
        (b"PKbad", "file.xlsx", "INVALID_XLSX"),
        (b"x", "../file.csv", "INVALID_FILENAME"),
    ],
)
def test_invalid_sources_have_stable_error_codes(
    source: bytes, filename: str, code: str
) -> None:
    with pytest.raises(IngestionError) as error:
        ingest_document(source, filename=filename)
    assert error.value.code == code


def test_mime_mismatch_is_rejected() -> None:
    with pytest.raises(IngestionError) as error:
        ingest_document(b"a,b\n", filename="data.csv", mime_type="application/pdf")
    assert error.value.code == "MIME_MISMATCH"
