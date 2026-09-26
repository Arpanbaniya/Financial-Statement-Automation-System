"""Statement detection should be explainable and conservative."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from finance import StatementRule, detect_statements
from ingestion import ingest_document
from ingestion.types import ExtractedDocument, ExtractedPage, IngestionWarning


def pdf_document(*page_texts: str) -> ExtractedDocument:
    return ExtractedDocument(
        source_filename="report.pdf",
        file_type="pdf",
        pages=tuple(
            ExtractedPage(number, text)
            for number, text in enumerate(page_texts, start=1)
        ),
        sheets=(),
        raw_text="\n\n".join(page_texts),
        tables=(),
        warnings=(),
        extraction_method="fixture",
        extraction_version="1",
    )


def test_multipage_pdf_detects_three_types_and_unknown_with_page_evidence() -> None:
    document = pdf_document(
        "ACME\nCONSOLIDATED STATEMENTS OF OPERATIONS\nRevenue 100\n"
        "Cost of revenue 40\nNet income 30",
        "CONSOLIDATED BALANCE SHEETS\nTotal assets 100\n"
        "Total liabilities 50\nStockholders' equity 50",
        "STATEMENT OF CASH FLOWS\nCash flows from operating activities\n"
        "Net cash provided by operating activities\n"
        "Cash flows from investing activities",
        "Management discussion\nRevenue increased during the year",
    )
    result = detect_statements(document)

    assert [item.statement_type for item in result] == [
        "income_statement",
        "balance_sheet",
        "cash_flow_statement",
        "unknown",
    ]
    assert [item.source_page for item in result] == [1, 2, 3, 4]
    assert all(0 < item.confidence < 1 for item in result[:3])
    assert result[0].evidence[0].source_reference == "page 1 line 2"
    assert result[0].evidence[0].kind == "heading"
    assert result[3].confidence == 0
    assert result[3].warnings[-1].code == "WEAK_STATEMENT_SIGNALS"


@pytest.mark.parametrize(
    ("heading", "expected"),
    [
        ("Statement of Income", "income_statement"),
        ("Profit and Loss", "income_statement"),
        ("P&L", "income_statement"),
        ("Statement of Financial Position", "balance_sheet"),
        ("Cash Flow Statement", "cash_flow_statement"),
    ],
)
def test_heading_naming_conventions(heading: str, expected: str) -> None:
    result = detect_statements(pdf_document(f"{heading}\n2025\n2024"))
    assert result[0].statement_type == expected
    assert result[0].confidence < 1


def test_spreadsheet_sheets_and_csv_rows_are_classified_by_line_items() -> None:
    workbook = Workbook()
    income = workbook.active
    income.title = "P&L"
    income["A1"] = "Revenue"
    income["A2"] = "Cost of sales"
    balance = workbook.create_sheet("Position")
    balance["A1"] = "Total assets"
    balance["A2"] = "Total liabilities"
    flows = workbook.create_sheet("Cash")
    flows["A1"] = "Cash flows from operating activities"
    flows["A2"] = "Cash flows from investing activities"
    output = BytesIO()
    workbook.save(output)
    document = ingest_document(output.getvalue(), filename="report.xlsx")

    result = detect_statements(document)
    assert [item.statement_type for item in result] == [
        "income_statement",
        "balance_sheet",
        "cash_flow_statement",
    ]
    assert [item.source_sheet for item in result] == ["P&L", "Position", "Cash"]
    assert result[1].evidence[0].source_reference == "Position!A1"

    csv_document = ingest_document(
        b"Item,2025\nRevenue,100\nCost of sales,40\nNet income,30\n",
        filename="figures.csv",
    )
    csv_result = detect_statements(csv_document)[0]
    assert csv_result.statement_type == "income_statement"
    assert csv_result.source_table == "CSV records"
    assert any(
        evidence.source_reference == "record 2 (A2)" for evidence in csv_result.evidence
    )


def test_separate_regions_on_one_sheet_do_not_mix_signals() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Overview"
    sheet["A1"] = "Revenue"
    sheet["A2"] = "Cost of sales"
    sheet["A5"] = "Total assets"
    sheet["A6"] = "Total liabilities"
    output = BytesIO()
    workbook.save(output)

    result = detect_statements(
        ingest_document(output.getvalue(), filename="combined.xlsx")
    )
    assert [item.statement_type for item in result] == [
        "income_statement",
        "balance_sheet",
    ]
    assert [item.source_table for item in result] == [
        "Overview region 1",
        "Overview region 2",
    ]


def test_sheet_heading_separated_by_blank_row_stays_with_its_data() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet["A1"] = "Statement of Operations"
    sheet["A3"] = "Revenue"
    sheet["A4"] = "Cost of sales"
    output = BytesIO()
    workbook.save(output)

    result = detect_statements(
        ingest_document(output.getvalue(), filename="spaced.xlsx")
    )
    assert len(result) == 1
    assert result[0].statement_type == "income_statement"
    assert {evidence.source_reference for evidence in result[0].evidence} >= {
        "Report!A1",
        "Report!A3",
    }


def test_empty_named_sheet_is_not_classified_from_its_name_alone() -> None:
    workbook = Workbook()
    workbook.active.title = "Balance Sheet"
    output = BytesIO()
    workbook.save(output)
    result = detect_statements(
        ingest_document(output.getvalue(), filename="empty.xlsx")
    )
    assert result[0].statement_type == "unknown"
    assert any(warning.code == "EMPTY_SHEET" for warning in result[0].warnings)


def test_conflicting_signals_and_scanned_pages_remain_unknown() -> None:
    document = pdf_document(
        "Total assets\nTotal liabilities\nRevenue\nCost of revenue",
        "",
    )
    document = ExtractedDocument(
        source_filename=document.source_filename,
        file_type=document.file_type,
        pages=document.pages,
        sheets=document.sheets,
        raw_text=document.raw_text,
        tables=document.tables,
        warnings=(IngestionWarning("OCR_UNSUPPORTED", "OCR unavailable", "page 2"),),
        extraction_method=document.extraction_method,
        extraction_version=document.extraction_version,
        extraction_status="needs_review",
    )
    result = detect_statements(document)
    assert result[0].statement_type == "unknown"
    assert any(warning.code == "AMBIGUOUS_STATEMENT" for warning in result[0].warnings)
    assert {item.statement_type for item in result[0].evidence} == {
        "income_statement",
        "balance_sheet",
    }
    assert result[1].statement_type == "unknown"
    assert any(warning.code == "OCR_UNSUPPORTED" for warning in result[1].warnings)


def test_rules_are_configurable_and_reject_duplicate_types() -> None:
    document = pdf_document("Earnings Report\nSales\nCosts")
    custom = StatementRule(
        "income_statement",
        (r"earnings report",),
        (r"sales\b", r"costs\b"),
    )
    result = detect_statements(document, rules=(custom,))
    assert result[0].statement_type == "income_statement"
    with pytest.raises(ValueError, match="unique"):
        detect_statements(document, rules=(custom, custom))
