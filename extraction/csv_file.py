"""Read CSV records while retaining decoded source text and row references."""

import csv
import sys
from io import StringIO

from openpyxl.utils import get_column_letter

from ingestion.errors import IngestionError
from ingestion.limits import MAX_CELLS, MAX_ROWS, MAX_TEXT_CHARS
from ingestion.types import (
    ExtractedCell,
    ExtractedDocument,
    ExtractedRow,
    ExtractedTable,
    IngestionWarning,
)


def _decode(source: bytes) -> tuple[str, str, list[IngestionWarning]]:
    warnings: list[IngestionWarning] = []
    try:
        if source.startswith((b"\xff\xfe", b"\xfe\xff")):
            text, encoding = source.decode("utf-16"), "utf-16"
        elif source.startswith(b"\xef\xbb\xbf"):
            text, encoding = source.decode("utf-8-sig"), "utf-8-sig"
        else:
            try:
                text, encoding = source.decode("utf-8"), "utf-8"
            except UnicodeDecodeError:
                text, encoding = source.decode("cp1252"), "cp1252"
                warnings.append(
                    IngestionWarning(
                        "LEGACY_ENCODING",
                        "Decoded as Windows-1252; verify accented characters",
                    )
                )
    except UnicodeDecodeError as exc:
        raise IngestionError(
            "UNSUPPORTED_ENCODING", "The CSV encoding could not be decoded safely"
        ) from exc
    if any(ord(char) < 32 and char not in "\r\n\t" for char in text) or "\x7f" in text:
        raise IngestionError(
            "UNSUPPORTED_ENCODING", "The CSV contains binary control characters"
        )
    if not text.strip():
        raise IngestionError("EMPTY_CSV", "The CSV has no records")
    if len(text) > MAX_TEXT_CHARS:
        raise IngestionError("DOCUMENT_TOO_COMPLEX", "Decoded CSV text is too large")
    return text, encoding, warnings


def extract_csv(source: bytes, filename: str) -> ExtractedDocument:
    text, encoding, warnings = _decode(source)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
        warnings.append(
            IngestionWarning(
                "DELIMITER_UNCERTAIN", "Could not detect a delimiter; assumed comma"
            )
        )

    rows: list[ExtractedRow] = []
    expected_columns: int | None = None
    cell_count = 0
    try:
        reader = csv.reader(
            StringIO(text, newline=""), delimiter=delimiter, strict=True
        )
        for record_number, fields in enumerate(reader, start=1):
            if record_number > MAX_ROWS:
                raise IngestionError("DOCUMENT_TOO_COMPLEX", "Too many CSV records")
            cell_count += len(fields)
            if cell_count > MAX_CELLS:
                raise IngestionError("DOCUMENT_TOO_COMPLEX", "Too many CSV cells")
            if expected_columns is None:
                expected_columns = len(fields)
            elif len(fields) != expected_columns:
                warnings.append(
                    IngestionWarning(
                        "COLUMN_COUNT_MISMATCH",
                        f"Found {len(fields)} columns; expected {expected_columns}",
                        f"record {record_number}",
                    )
                )
            rows.append(
                ExtractedRow(
                    number=record_number,
                    source_line=reader.line_num,
                    cells=tuple(
                        ExtractedCell(
                            coordinate=f"{get_column_letter(column)}{record_number}",
                            value=value,
                            data_type="text",
                        )
                        for column, value in enumerate(fields, start=1)
                    ),
                )
            )
    except csv.Error as exc:
        raise IngestionError(
            "INVALID_CSV", "The CSV records could not be read"
        ) from exc
    if not rows:
        raise IngestionError("EMPTY_CSV", "The CSV has no records")

    # The first record is kept in rows too. Its labels are only a candidate header.
    headers = tuple(str(cell.value) for cell in rows[0].cells)
    return ExtractedDocument(
        source_filename=filename,
        file_type="csv",
        pages=(),
        sheets=(),
        raw_text=text,
        tables=(
            ExtractedTable(
                name="CSV records",
                source="CSV",
                rows=tuple(rows),
                headers=headers,
                delimiter=delimiter,
            ),
        ),
        warnings=tuple(warnings),
        extraction_method="python-csv",
        extraction_version=sys.version.split()[0],
        source_encoding=encoding,
    )
