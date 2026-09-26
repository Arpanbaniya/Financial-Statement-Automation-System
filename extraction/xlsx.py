"""Read populated worksheet regions without changing or evaluating a workbook."""

from datetime import date, datetime, time, timedelta
from importlib.metadata import version
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

import ingestion.limits as limits
from ingestion.errors import IngestionError
from ingestion.types import (
    CellValue,
    ExtractedCell,
    ExtractedDocument,
    ExtractedRow,
    ExtractedSheet,
    ExtractedTable,
    IngestionWarning,
)


def _make_region(name: str, number: int, rows: list[ExtractedRow]) -> ExtractedTable:
    columns = [
        column_index_from_string(coordinate_from_string(cell.coordinate)[0])
        for row in rows
        for cell in row.cells
    ]
    reference = (
        f"{get_column_letter(min(columns))}{rows[0].number}:"
        f"{get_column_letter(max(columns))}{rows[-1].number}"
    )
    return ExtractedTable(
        name=f"{name} region {number}",
        source=name,
        rows=tuple(rows),
        range_reference=reference,
    )


def _populated_regions(
    name: str, rows: tuple[ExtractedRow, ...]
) -> list[ExtractedTable]:
    """Separate blocks divided by one or more fully blank worksheet rows."""
    regions: list[ExtractedTable] = []
    block: list[ExtractedRow] = []
    for row in rows:
        if block and row.number > block[-1].number + 1:
            regions.append(_make_region(name, len(regions) + 1, block))
            block = []
        block.append(row)
    if block:
        regions.append(_make_region(name, len(regions) + 1, block))
    return regions


def extract_xlsx(source: bytes, filename: str) -> ExtractedDocument:
    if not source.startswith(b"PK"):
        raise IngestionError("INVALID_XLSX", "This is not an XLSX workbook")
    try:
        with ZipFile(BytesIO(source)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > limits.MAX_ZIP_ENTRIES
                or sum(entry.file_size for entry in entries)
                > limits.MAX_UNCOMPRESSED_XLSX_BYTES
            ):
                raise IngestionError(
                    "DOCUMENT_TOO_COMPLEX", "The workbook expands beyond safe limits"
                )
        workbook = load_workbook(
            BytesIO(source), read_only=True, data_only=False, keep_links=False
        )
    except IngestionError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, TypeError) as exc:
        raise IngestionError("INVALID_XLSX", "The workbook could not be read") from exc

    sheets: list[ExtractedSheet] = []
    tables: list[ExtractedTable] = []
    text_parts: list[str] = []
    warnings: list[IngestionWarning] = []
    row_count = cell_count = text_chars = 0
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            rows: list[ExtractedRow] = []
            text_parts.append(f"[{worksheet.title}]")
            for row in worksheet.iter_rows():
                row_count += 1
                if row_count > limits.MAX_ROWS:
                    raise IngestionError(
                        "DOCUMENT_TOO_COMPLEX", "Too many worksheet rows"
                    )
                cells: list[ExtractedCell] = []
                for cell in row:
                    if cell.value is None:
                        continue
                    cell_count += 1
                    if cell_count > limits.MAX_CELLS:
                        raise IngestionError(
                            "DOCUMENT_TOO_COMPLEX", "Too many workbook cells"
                        )
                    value: CellValue = cell.value
                    if not isinstance(
                        value, (str, int, float, bool, date, datetime, time, timedelta)
                    ):
                        value = str(value)
                        warnings.append(
                            IngestionWarning(
                                "CELL_TYPE_COERCED",
                                "An uncommon cell value was preserved as text",
                                f"{worksheet.title}!{cell.coordinate}",
                            )
                        )
                    cells.append(
                        ExtractedCell(
                            coordinate=cell.coordinate,
                            value=value,
                            data_type=cell.data_type,
                        )
                    )
                if cells:
                    source_row = next(
                        cell.row for cell in row if cell.value is not None
                    )
                    rows.append(ExtractedRow(number=source_row, cells=tuple(cells)))
                    rendered = "\t".join(str(cell.value) for cell in cells)
                    text_chars += len(rendered)
                    if text_chars > limits.MAX_TEXT_CHARS:
                        raise IngestionError(
                            "DOCUMENT_TOO_COMPLEX", "Extracted text is too large"
                        )
                    text_parts.append(rendered)
            frozen_rows = tuple(rows)
            sheets.append(
                ExtractedSheet(
                    name=worksheet.title, index=sheet_index, rows=frozen_rows
                )
            )
            tables.extend(_populated_regions(worksheet.title, frozen_rows))
            if not frozen_rows:
                warnings.append(
                    IngestionWarning(
                        "EMPTY_SHEET", "No populated cells were found", worksheet.title
                    )
                )
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("INVALID_XLSX", "The workbook could not be read") from exc
    finally:
        workbook.close()

    return ExtractedDocument(
        source_filename=filename,
        file_type="xlsx",
        pages=(),
        sheets=tuple(sheets),
        raw_text="\n".join(text_parts),
        tables=tuple(tables),
        warnings=tuple(warnings),
        extraction_method="openpyxl-read-only",
        extraction_version=version("openpyxl"),
    )
