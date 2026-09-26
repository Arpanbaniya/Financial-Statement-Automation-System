"""Page text and candidate table extraction from PDF source bytes."""

from importlib.metadata import version
from io import BytesIO

import pdfplumber
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from ingestion.errors import IngestionError
from ingestion.limits import MAX_CELLS, MAX_PDF_PAGES, MAX_ROWS, MAX_TEXT_CHARS
from ingestion.types import (
    ExtractedCell,
    ExtractedDocument,
    ExtractedPage,
    ExtractedRow,
    ExtractedTable,
    IngestionWarning,
)


def extract_pdf(source: bytes, filename: str) -> ExtractedDocument:
    if not source.startswith(b"%PDF-"):
        raise IngestionError("INVALID_PDF", "This is not a PDF file")
    try:
        reader = PdfReader(BytesIO(source), strict=False)
        if reader.is_encrypted:
            raise IngestionError(
                "ENCRYPTED_PDF", "Password-protected PDFs are not supported"
            )
        page_count = len(reader.pages)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("INVALID_PDF", "The PDF could not be read") from exc
    if page_count > MAX_PDF_PAGES:
        raise IngestionError("DOCUMENT_TOO_COMPLEX", "The PDF has too many pages")
    if page_count == 0:
        raise IngestionError("EMPTY_PDF", "The PDF has no pages")

    warnings: list[IngestionWarning] = []
    pages: list[ExtractedPage] = []
    tables: list[ExtractedTable] = []
    text_chars = row_count = cell_count = 0
    try:
        plumber = pdfplumber.open(BytesIO(source))
    except Exception:
        plumber = None
        warnings.append(
            IngestionWarning(
                "PDF_TABLE_READER_UNAVAILABLE",
                "The PDF table reader could not open this file",
            )
        )

    try:
        for number in range(1, page_count + 1):
            text = ""
            method = "pypdf-text"
            try:
                text = reader.pages[number - 1].extract_text() or ""
            except Exception:
                warnings.append(
                    IngestionWarning(
                        "PAGE_TEXT_FAILED",
                        "Primary text extraction failed",
                        f"page {number}",
                    )
                )

            plumber_page = None
            if plumber is not None:
                try:
                    plumber_page = plumber.pages[number - 1]
                except Exception:
                    warnings.append(
                        IngestionWarning(
                            "PDF_PAGE_READER_FAILED",
                            "The table reader could not open this page",
                            f"page {number}",
                        )
                    )

            if not text.strip() and plumber_page is not None:
                try:
                    fallback = plumber_page.extract_text() or ""
                    if fallback.strip():
                        text = fallback
                        method = "pdfplumber-text"
                except Exception:
                    warnings.append(
                        IngestionWarning(
                            "PAGE_TEXT_FALLBACK_FAILED",
                            "Alternate text extraction failed",
                            f"page {number}",
                        )
                    )
            if not text.strip():
                warnings.append(
                    IngestionWarning(
                        "NO_EXTRACTABLE_TEXT",
                        "No searchable text was found on this page",
                        f"page {number}",
                    )
                )
            text_chars += len(text)
            if text_chars > MAX_TEXT_CHARS:
                raise IngestionError(
                    "DOCUMENT_TOO_COMPLEX", "Extracted text is too large"
                )
            pages.append(ExtractedPage(number=number, text=text, method=method))

            if plumber_page is None:
                continue
            try:
                candidates = plumber_page.extract_tables()
            except Exception:
                warnings.append(
                    IngestionWarning(
                        "TABLE_EXTRACTION_FAILED",
                        "Table detection failed on this page",
                        f"page {number}",
                    )
                )
                continue
            for table_number, candidate in enumerate(candidates, start=1):
                rows: list[ExtractedRow] = []
                for record_number, fields in enumerate(candidate, start=1):
                    row_count += 1
                    cell_count += len(fields)
                    if row_count > MAX_ROWS or cell_count > MAX_CELLS:
                        raise IngestionError(
                            "DOCUMENT_TOO_COMPLEX", "Too many extracted PDF table cells"
                        )
                    rows.append(
                        ExtractedRow(
                            number=record_number,
                            cells=tuple(
                                ExtractedCell(
                                    coordinate=f"{get_column_letter(column)}{record_number}",
                                    value=value,
                                    data_type="text",
                                )
                                for column, value in enumerate(fields, start=1)
                            ),
                        )
                    )
                if rows:
                    tables.append(
                        ExtractedTable(
                            name=f"Page {number} table {table_number}",
                            source=f"page {number}",
                            rows=tuple(rows),
                            page_number=number,
                        )
                    )
    finally:
        if plumber is not None:
            plumber.close()

    needs_review = not any(page.text.strip() for page in pages)
    if needs_review:
        warnings.append(
            IngestionWarning(
                "OCR_UNSUPPORTED",
                "No searchable text. OCR is unsupported; review the source manually.",
            )
        )
    return ExtractedDocument(
        source_filename=filename,
        file_type="pdf",
        pages=tuple(pages),
        sheets=(),
        raw_text="\n\n".join(page.text for page in pages),
        tables=tuple(tables),
        warnings=tuple(warnings),
        extraction_method="pypdf+pdfplumber",
        extraction_version=(
            f"pypdf {version('pypdf')}; pdfplumber {version('pdfplumber')}"
        ),
        extraction_status="needs_review" if needs_review else "extracted",
    )
