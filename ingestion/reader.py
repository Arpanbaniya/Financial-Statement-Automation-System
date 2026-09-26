"""Small, bounded readers for untrusted PDF, XLSX, and CSV source bytes."""

import csv
import sys
from datetime import date, datetime, time, timedelta
from importlib.metadata import version
from io import BytesIO, StringIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from ingestion.types import (
    CellValue,
    ExtractedCell,
    ExtractedDocument,
    ExtractedPage,
    ExtractedRow,
    ExtractedSheet,
    ExtractedTable,
    FileType,
    IngestionWarning,
)

MAX_SOURCE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_XLSX_BYTES = 60 * 1024 * 1024
MAX_ZIP_ENTRIES = 2_000
MAX_PDF_PAGES = 500
MAX_ROWS = 50_000
MAX_CELLS = 250_000
MAX_TEXT_CHARS = 10_000_000


class IngestionError(ValueError):
    """A source cannot be ingested; code is safe to record in a job."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


def ingest_document(
    source: bytes, *, filename: str, mime_type: str | None = None
) -> ExtractedDocument:
    """Return source content and warnings without changing or interpreting the file.

    Accepts bytes downloaded from the private bucket by a future processing step.
    Unsupported or irrecoverably malformed files raise a coded IngestionError.
    """
    if not isinstance(source, bytes):
        raise IngestionError("INVALID_SOURCE", "Source must be bytes")
    if not source:
        raise IngestionError("EMPTY_SOURCE", "The source file is empty")
    if len(source) > MAX_SOURCE_BYTES:
        raise IngestionError("SOURCE_TOO_LARGE", "The source exceeds 10 MB")
    if (
        not filename
        or len(filename) > 255
        or "/" in filename
        or "\\" in filename
        or any(ord(char) < 32 or ord(char) == 127 for char in filename)
    ):
        raise IngestionError("INVALID_FILENAME", "A plain source filename is required")

    extension = PurePath(filename).suffix.lower()
    file_type: FileType
    if extension == ".pdf":
        file_type = "pdf"
        allowed_mimes = {"application/pdf"}
    elif extension == ".xlsx":
        file_type = "xlsx"
        allowed_mimes = {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }
    elif extension == ".csv":
        file_type = "csv"
        allowed_mimes = {"text/csv", "application/csv", "application/vnd.ms-excel"}
    else:
        raise IngestionError(
            "UNSUPPORTED_FILE_TYPE", "Only PDF, XLSX, and CSV are supported"
        )
    if mime_type is not None and mime_type not in allowed_mimes:
        raise IngestionError("MIME_MISMATCH", "MIME type does not match the filename")

    if file_type == "pdf":
        return _read_pdf(source, filename)
    if file_type == "xlsx":
        return _read_xlsx(source, filename)
    return _read_csv(source, filename)


def _read_pdf(source: bytes, filename: str) -> ExtractedDocument:
    if not source.startswith(b"%PDF-"):
        raise IngestionError("INVALID_PDF", "This is not a PDF file")
    try:
        reader = PdfReader(BytesIO(source), strict=False)
        if reader.is_encrypted:
            raise IngestionError(
                "ENCRYPTED_PDF", "Password-protected PDFs are not supported"
            )
        page_count = len(reader.pages)
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("INVALID_PDF", "The PDF could not be read") from exc
    if page_count > MAX_PDF_PAGES:
        raise IngestionError("DOCUMENT_TOO_COMPLEX", "The PDF has too many pages")
    if page_count == 0:
        raise IngestionError("EMPTY_PDF", "The PDF has no pages")

    pages: list[ExtractedPage] = []
    warnings: list[IngestionWarning] = []
    total_chars = 0
    for number in range(1, page_count + 1):
        try:
            page = reader.pages[number - 1]
            text = page.extract_text() or ""
        except Exception:
            text = ""
            warnings.append(
                IngestionWarning(
                    "PAGE_TEXT_FAILED",
                    "Text extraction failed on this page",
                    f"page {number}",
                )
            )
        if not text.strip():
            warnings.append(
                IngestionWarning(
                    "NO_EXTRACTABLE_TEXT",
                    "No searchable text was found; scanned images need OCR",
                    f"page {number}",
                )
            )
        total_chars += len(text)
        if total_chars > MAX_TEXT_CHARS:
            raise IngestionError("DOCUMENT_TOO_COMPLEX", "Extracted text is too large")
        pages.append(ExtractedPage(number=number, text=text))

    return ExtractedDocument(
        source_filename=filename,
        file_type="pdf",
        pages=tuple(pages),
        sheets=(),
        raw_text="\n\n".join(page.text for page in pages),
        tables=(),
        warnings=tuple(warnings),
        extraction_method="pypdf-text",
        extraction_version=version("pypdf"),
    )


def _read_xlsx(source: bytes, filename: str) -> ExtractedDocument:
    if not source.startswith(b"PK"):
        raise IngestionError("INVALID_XLSX", "This is not an XLSX workbook")
    try:
        with ZipFile(BytesIO(source)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > MAX_ZIP_ENTRIES
                or sum(entry.file_size for entry in entries)
                > MAX_UNCOMPRESSED_XLSX_BYTES
            ):
                raise IngestionError(
                    "DOCUMENT_TOO_COMPLEX", "The workbook expands beyond safe limits"
                )
        workbook = load_workbook(
            BytesIO(source), read_only=True, data_only=False, keep_links=False
        )
    except IngestionError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError, TypeError) as exc:
        raise IngestionError("INVALID_XLSX", "The workbook could not be read") from exc

    sheets: list[ExtractedSheet] = []
    tables: list[ExtractedTable] = []
    text_parts: list[str] = []
    warnings: list[IngestionWarning] = []
    row_count = cell_count = text_chars = 0
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            rows: list[ExtractedRow] = []
            text_parts.append(f"[{worksheet.title}]")
            for row in worksheet.iter_rows():
                row_count += 1
                if row_count > MAX_ROWS:
                    raise IngestionError(
                        "DOCUMENT_TOO_COMPLEX", "Too many worksheet rows"
                    )
                cells: list[ExtractedCell] = []
                for cell in row:
                    if cell.value is None:
                        continue
                    cell_count += 1
                    if cell_count > MAX_CELLS:
                        raise IngestionError(
                            "DOCUMENT_TOO_COMPLEX", "Too many workbook cells"
                        )
                    value: CellValue = cell.value
                    if not isinstance(
                        value, (str, int, float, bool, date, datetime, time, timedelta)
                    ):
                        value = str(value)
                        warnings.append(
                            IngestionWarning(
                                "CELL_TYPE_COERCED",
                                "An uncommon cell value was preserved as text",
                                f"{worksheet.title}!{cell.coordinate}",
                            )
                        )
                    cells.append(
                        ExtractedCell(
                            coordinate=cell.coordinate,
                            value=value,
                            data_type=cell.data_type,
                        )
                    )
                if cells:
                    source_row = next(
                        cell.row for cell in row if cell.value is not None
                    )
                    rows.append(ExtractedRow(number=source_row, cells=tuple(cells)))
                    rendered = "\t".join(str(cell.value) for cell in cells)
                    text_chars += len(rendered)
                    if text_chars > MAX_TEXT_CHARS:
                        raise IngestionError(
                            "DOCUMENT_TOO_COMPLEX", "Extracted text is too large"
                        )
                    text_parts.append(rendered)
            frozen_rows = tuple(rows)
            sheets.append(
                ExtractedSheet(
                    name=worksheet.title, index=sheet_index, rows=frozen_rows
                )
            )
            tables.append(
                ExtractedTable(
                    name=worksheet.title,
                    source=worksheet.title,
                    rows=frozen_rows,
                )
            )
    except IngestionError:
        raise
    except Exception as exc:
        raise IngestionError("INVALID_XLSX", "The workbook could not be read") from exc
    finally:
        workbook.close()

    return ExtractedDocument(
        source_filename=filename,
        file_type="xlsx",
        pages=(),
        sheets=tuple(sheets),
        raw_text="\n".join(text_parts),
        tables=tuple(tables),
        warnings=tuple(warnings),
        extraction_method="openpyxl-read-only",
        extraction_version=version("openpyxl"),
    )


def _read_csv(source: bytes, filename: str) -> ExtractedDocument:
    try:
        raw_text = source.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise IngestionError("UNSUPPORTED_ENCODING", "CSV must use UTF-8") from exc
    if not raw_text.strip():
        raise IngestionError("EMPTY_CSV", "The CSV has no records")
    warnings: list[IngestionWarning] = []
    try:
        dialect = csv.Sniffer().sniff(raw_text[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
        warnings.append(
            IngestionWarning(
                "DELIMITER_ASSUMED", "The delimiter was unclear; comma was used"
            )
        )

    rows: list[ExtractedRow] = []
    cell_count = 0
    try:
        reader = csv.reader(
            StringIO(raw_text, newline=""), delimiter=delimiter, strict=True
        )
        expected_columns: int | None = None
        for record_number, fields in enumerate(reader, start=1):
            if record_number > MAX_ROWS:
                raise IngestionError("DOCUMENT_TOO_COMPLEX", "Too many CSV records")
            cell_count += len(fields)
            if cell_count > MAX_CELLS:
                raise IngestionError("DOCUMENT_TOO_COMPLEX", "Too many CSV fields")
            if expected_columns is None:
                expected_columns = len(fields)
            elif len(fields) != expected_columns:
                warnings.append(
                    IngestionWarning(
                        "COLUMN_COUNT_MISMATCH",
                        f"Expected {expected_columns} columns, found {len(fields)}",
                        f"record {record_number}",
                    )
                )
            cells = tuple(
                ExtractedCell(
                    coordinate=f"{get_column_letter(column)}{record_number}",
                    value=value,
                    data_type="text",
                )
                for column, value in enumerate(fields, start=1)
            )
            rows.append(
                ExtractedRow(
                    number=record_number, cells=cells, source_line=reader.line_num
                )
            )
    except csv.Error as exc:
        raise IngestionError("INVALID_CSV", "The CSV could not be read") from exc
    if not rows:
        raise IngestionError("EMPTY_CSV", "The CSV has no records")

    return ExtractedDocument(
        source_filename=filename,
        file_type="csv",
        pages=(),
        sheets=(),
        raw_text=raw_text,
        tables=(
            ExtractedTable(
                name=filename, source="CSV", rows=tuple(rows), delimiter=delimiter
            ),
        ),
        warnings=tuple(warnings),
        extraction_method="python-csv",
        extraction_version=f"{sys.version_info.major}.{sys.version_info.minor}",
    )
