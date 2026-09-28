"""Trial-balance audit schedules added to the existing snapshot workbook."""

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

from reports.excel import _safe_text, build_excel_report


def build_trial_balance_report(
    company, result, request, statements, ratios, working, flows, checks
):
    output = build_excel_report(
        company_name=company,
        statements=tuple(
            item for item in statements if item.period.period_end == request.period_end
        ),
        ratios=ratios,
        working_capital=working,
        cash_flow=flows,
        validations=checks,
        generated_at=datetime.now(UTC),
    )
    workbook = load_workbook(BytesIO(output))
    mapping = workbook.create_sheet("Trial Balance")
    mapping.append(
        [
            "Source row",
            "Code",
            "Account",
            "Debit",
            "Credit",
            "Opening debit",
            "Opening credit",
            "Category",
            "Included",
        ]
    )
    for row in result["rows"]:
        mapping.append(
            [
                row["id"],
                row["code"],
                row["account"],
                *[
                    float(Decimal(row[field]))
                    for field in ("debit", "credit", "opening_debit", "opening_credit")
                ],
                row["category"] or "Unmapped",
                "No" if row["excluded"] else "Yes",
            ]
        )
        for col in (1, 2, 3, 8, 9):
            _safe_text(
                mapping.cell(mapping.max_row, col),
                mapping.cell(mapping.max_row, col).value,
            )
    schedule = workbook.create_sheet("Generation Checks")
    schedule.append(["Check / input", "Value", "Explanation"])
    schedule.append(
        [
            "Reporting basis",
            request.basis,
            "Current pre-closing profit is added to retained earnings once.",
        ]
    )
    schedule.append(
        [
            "Currency",
            request.currency,
            f"Imported amounts: {request.unit_scale}; statements: ones.",
        ]
    )
    schedule.append(
        [
            "Debit total",
            float(Decimal(result["debit"])),
            "Included source accounts only",
        ]
    )
    schedule.append(
        [
            "Credit total",
            float(Decimal(result["credit"])),
            "Included source accounts only",
        ]
    )
    schedule.append(
        [
            "Difference",
            float(Decimal(result["difference"])),
            "Maximum tolerance: 0.01 normalized currency units",
        ]
    )
    for name, value in result["retained_earnings"].items():
        schedule.append(
            [
                "Retained earnings: " + name.replace("_", " "),
                float(value),
                "Brought forward + profit - distributions = closing",
            ]
        )
    for name, value in result.get("cash_bridge", {}).items():
        schedule.append(
            [
                name.replace("_", " ").title(),
                float(value),
                "Cash bridge / operating reconciliation in normalized ones",
            ]
        )
    for issue in result["cash_flow_issues"]:
        schedule.append(["Cash flow unavailable", None, issue])
    if request.cash:
        for name, value in request.cash.model_dump().items():
            schedule.append(
                [
                    name.replace("_", " ").title(),
                    float(value) if isinstance(value, Decimal) else str(value),
                    "User-confirmed cash schedule; amounts in imported units",
                ]
            )
    for row in schedule:
        for cell in row:
            if isinstance(cell.value, str):
                _safe_text(cell, cell.value)
    for sheet in (mapping, schedule):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="173F32")
            cell.font = Font(color="FFFFFF", bold=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                65, max(18, max(len(str(c.value or "")) for c in column) + 2)
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00;[Red](#,##0.00);"—"'
    chart_sheet = workbook.create_sheet("Charts")
    chart_sheet.append(["Profitability", request.currency])
    income = result["statements"].get("income_statement", {})
    for name in ("revenue", "gross_profit", "operating_income", "net_income"):
        chart_sheet.append(
            [
                name.replace("_", " ").title(),
                float(income[name]) if name in income else None,
            ]
        )
    chart_sheet.append([])
    chart_sheet.append(["Financial position", request.currency])
    for name in ("total_assets", "total_liabilities", "shareholders_equity"):
        chart_sheet.append(
            [
                name.replace("_", " ").title(),
                float(result["statements"]["balance_sheet"][name]),
            ]
        )
    for title, start, end, anchor in (
        ("Profitability", 1, 5, "D2"),
        ("Financial position", 7, 10, "D18"),
    ):
        chart = BarChart()
        chart.title = title
        chart.y_axis.title = request.currency
        chart.add_data(
            Reference(chart_sheet, min_col=2, min_row=start, max_row=end),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(chart_sheet, min_col=1, min_row=start + 1, max_row=end)
        )
        chart.height, chart.width = 8, 17
        chart_sheet.add_chart(chart, anchor)
    chart_sheet.column_dimensions["A"].width = 30
    chart_sheet.column_dimensions["B"].width = 22
    workbook["Methodology"].append(
        [
            "Trial balance generation",
            "Snapshot export. Re-upload and regenerate after editing source amounts. "
            "Trial Balance records mappings and exclusions; Generation Checks "
            "records retained earnings and cash inputs.",
        ]
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
