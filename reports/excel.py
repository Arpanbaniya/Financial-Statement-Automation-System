"""Generate a transparent, source-linked financial analysis workbook."""

import re
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from finance.cash_flow import CashFlowResult
from finance.commentary import CommentaryFinding
from finance.common_size import CommonSizeResult
from finance.horizontal import HorizontalResult
from finance.ratio_registry import FORMULAS
from finance.ratios import MetricResult
from finance.working_capital import WorkingCapitalResult
from normalization.taxonomy import get_field
from validation.types import SourceRef, StatementSnapshot, ValidationResult

SHEETS = (
    "Summary",
    "Income Statement",
    "Balance Sheet",
    "Cash Flow",
    "Ratios",
    "Horizontal Analysis",
    "Common Size",
    "Working Capital",
    "Validation",
    "Source Data",
    "Methodology",
)
_NAVY = "173F32"
_NUMBER = "#,##0.00;[Red](#,##0.00);0.00"
_PERCENT = '0.00"%";[Red](0.00"%");0.00"%"'


def report_filename(company_name: str, period_end: str) -> str:
    """Use a safe, stable filename without leaking path characters."""
    company = re.sub(r"[^A-Za-z0-9_-]+", "_", company_name).strip("_")[:60]
    if not company:
        company = "Company"
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", period_end):
        raise ValueError("Period end must be YYYY-MM-DD")
    return f"Financial_Analysis_{company}_{period_end}.xlsx"


def _safe_text(cell, value: object) -> None:
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"  # Untrusted labels and original values cannot become formulas.


def _value(cell, value: Decimal | int | None, *, percent: bool = False) -> None:
    cell.value = float(value) if isinstance(value, Decimal) else value
    cell.number_format = _PERCENT if percent else _NUMBER


def _refs(refs: tuple[SourceRef, ...]) -> str:
    labels = []
    for ref in refs:
        location = ref.location
        position = (
            f"page {location.page}"
            if location.page
            else f"{location.sheet}!{location.cell}"
            if location.sheet and location.cell
            else f"{location.sheet} row {location.row}"
            if location.sheet and location.row
            else f"cell {location.cell}"
            if location.cell
            else f"line {location.line}"
            if location.line
            else "location unspecified"
        )
        labels.append(f"{ref.statement_id} / {ref.line_item_id or '-'} / {position}")
    return "; ".join(dict.fromkeys(labels))


def _sheet(
    wb: Workbook, name: str, title: str, subtitle: str, headers: tuple[str, ...]
):
    ws = wb[name]
    ws.sheet_view.showGridLines = False
    _safe_text(ws.cell(1, 1), title)
    ws.cell(1, 1).font = Font(size=17, bold=True, color=_NAVY)
    _safe_text(ws.cell(2, 1), subtitle)
    ws.cell(2, 1).font = Font(size=10, color="526C61")
    for column, header in enumerate(headers, 1):
        cell = ws.cell(4, column)
        _safe_text(cell, header)
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}4"
    return ws


def _append(
    ws,
    cells: tuple[object, ...],
    *,
    numeric: tuple[int, ...] = (),
    percent: tuple[int, ...] = (),
) -> None:
    row = ws.max_row + 1
    for column, value in enumerate(cells, 1):
        cell = ws.cell(row, column)
        if column in numeric or column in percent:
            _value(cell, value, percent=column in percent)
        elif isinstance(value, date):
            cell.value = value
            cell.number_format = "yyyy-mm-dd"
        else:
            _safe_text(cell, value)
        cell.alignment = Alignment(
            vertical="top", wrap_text=column > 2 or len(str(value or "")) > 48
        )
        if row % 2:
            cell.fill = PatternFill("solid", fgColor="F6F9F7")


def _finish(wb: Workbook) -> None:
    for ws in wb.worksheets:
        for column in ws.columns:
            letter = column[0].column_letter
            longest = max(len(str(cell.value or "")) for cell in column[:1000])
            ws.column_dimensions[letter].width = min(max(longest + 2, 13), 48)
        if ws.max_row > 4:
            ws.auto_filter.ref = f"A4:{get_column_letter(ws.max_column)}{ws.max_row}"
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0


def build_excel_report(
    *,
    company_name: str,
    statements: tuple[StatementSnapshot, ...],
    ratios: tuple[MetricResult, ...] = (),
    horizontal: tuple[HorizontalResult, ...] = (),
    common_size: tuple[CommonSizeResult, ...] = (),
    working_capital: tuple[WorkingCapitalResult, ...] = (),
    cash_flow: tuple[CashFlowResult, ...] = (),
    commentary: tuple[CommentaryFinding, ...] = (),
    validations: tuple[ValidationResult, ...] = (),
    generated_at: datetime,
) -> bytes:
    """Build the eleven required sheets from reviewed, source-linked records.

    Calculated values are snapshots of the Python engine, not editable Excel
    formulas. Source Data and Methodology explain every figure's derivation.
    """
    if generated_at.tzinfo is None or generated_at.utcoffset() is None:
        raise ValueError("Generation time must include a timezone")
    if not company_name.strip():
        raise ValueError("Company name is required")
    if not statements:
        raise ValueError("At least one statement is required")
    if len({item.company_id for item in statements}) != 1:
        raise ValueError("Statements must belong to one company")
    period_end = max(
        (item.period.period_end for item in statements if item.period.period_end),
        default=None,
    )
    if period_end is None:
        raise ValueError("A normalized period end is required")
    wb = Workbook()
    wb.active.title = SHEETS[0]
    for name in SHEETS[1:]:
        wb.create_sheet(name)
    wb.properties.title = f"Financial analysis — {company_name}"
    wb.properties.creator = "Financial Statement Automation System"
    wb.properties.created = generated_at.astimezone(UTC).replace(tzinfo=None)
    subtitle = (
        f"{company_name} · through {period_end.isoformat()} · "
        f"generated {generated_at.astimezone(UTC).isoformat()} · "
        "values in normalized ones"
    )
    currencies = {item.currency for item in statements if item.currency}
    currency_label = (
        next(iter(currencies)) if len(currencies) == 1 else "mixed currency"
    )
    summary = _sheet(
        wb,
        "Summary",
        "Financial analysis",
        subtitle,
        ("Measure", "Value", "Unit", "Status", "Source"),
    )
    for metric in (*cash_flow, *working_capital, *ratios):
        _append(
            summary,
            (
                metric.metric_name.replace("_", " ").title(),
                metric.value,
                currency_label if metric.unit == "currency" else metric.unit,
                metric.status,
                _refs(metric.source_refs),
            ),
            numeric=(2,),
        )
    if summary.max_row == 4:
        _append(
            summary, ("No calculated measures available", None, "", "unavailable", "")
        )
    if commentary:
        _append(summary, ("Factual observations", None, "", "", ""))
        for finding in commentary:
            _append(
                summary,
                (finding.text, None, "", finding.kind, _refs(finding.source_refs)),
            )
            summary.row_dimensions[summary.max_row].height = 45

    for statement_type, sheet_name in (
        ("income_statement", "Income Statement"),
        ("balance_sheet", "Balance Sheet"),
        ("cash_flow_statement", "Cash Flow"),
    ):
        ws = _sheet(
            wb,
            sheet_name,
            sheet_name,
            subtitle,
            (
                "Period end",
                "Field",
                "Accepted amount",
                "Currency",
                "Unit",
                "Statement ID",
                "Source",
            ),
        )
        for statement in statements:
            if statement.statement_type != statement_type:
                continue
            for item in statement.values:
                if item.status != "accepted" or item.normalized_value is None:
                    continue
                _append(
                    ws,
                    (
                        statement.period.period_end,
                        get_field(statement_type, item.field).display_name,
                        item.normalized_value,
                        item.currency or statement.currency,
                        "ones",
                        statement.id,
                        _refs((item.source_ref,)),
                    ),
                    numeric=(3,),
                )
        if ws.max_row == 4:
            _append(ws, ("No accepted values available", "", None, "", "", "", ""))

    ratio_ws = _sheet(
        wb,
        "Ratios",
        "Financial ratios",
        subtitle,
        (
            "Metric",
            "Value",
            "Unit",
            "Formula ID",
            "Numerator",
            "Denominator",
            "Status",
            "Warnings",
            "Source",
        ),
    )
    for item in ratios:
        _append(
            ratio_ws,
            (
                item.metric_name,
                item.value,
                currency_label if item.unit == "currency" else item.unit,
                item.formula_id,
                item.numerator,
                item.denominator,
                item.status,
                "; ".join(w.message for w in item.warnings),
                _refs(item.source_refs),
            ),
            numeric=(2, 5, 6),
        )

    growth_ws = _sheet(
        wb,
        "Horizontal Analysis",
        "Changes over time",
        subtitle,
        (
            "Metric",
            "Earlier",
            "Current",
            "Amount change",
            "Percent change",
            "Status",
            "Source",
        ),
    )
    for item in horizontal:
        _append(
            growth_ws,
            (
                item.metric_name,
                item.previous_value,
                item.current_value,
                item.absolute_change,
                item.percentage_change,
                item.status,
                _refs(item.source_refs),
            ),
            numeric=(2, 3, 4),
            percent=(5,),
        )

    common_ws = _sheet(
        wb,
        "Common Size",
        "Common-size statements",
        subtitle,
        (
            "Statement",
            "Field",
            "Amount",
            "Base",
            "Percent",
            "Status",
            "Source",
        ),
    )
    for item in common_size:
        _append(
            common_ws,
            (
                item.period.statement_type,
                item.field,
                item.amount,
                item.base_amount,
                item.percentage,
                item.status,
                _refs(item.source_refs),
            ),
            numeric=(3, 4),
            percent=(5,),
        )

    wc_ws = _sheet(
        wb,
        "Working Capital",
        "Working capital",
        subtitle,
        (
            "Metric",
            "Value",
            "Unit",
            "Days in period",
            "Formula ID",
            "Status",
            "Warnings",
            "Source",
        ),
    )
    for item in working_capital:
        _append(
            wc_ws,
            (
                item.metric_name,
                item.value,
                currency_label if item.unit == "currency" else item.unit,
                item.days_in_period,
                item.formula_id,
                item.status,
                "; ".join(w.message for w in item.warnings),
                _refs(item.source_refs),
            ),
            numeric=(2, 4),
        )

    validation_ws = _sheet(
        wb,
        "Validation",
        "Validation checks",
        subtitle,
        (
            "Check",
            "Status",
            "Severity",
            "Expected",
            "Actual",
            "Difference",
            "Explanation",
            "Source",
        ),
    )
    for item in validations:
        _append(
            validation_ws,
            (
                item.check_name,
                item.status,
                item.severity,
                item.expected_value,
                item.actual_value,
                item.difference,
                item.explanation,
                _refs(item.source_refs),
            ),
            numeric=(4, 5, 6),
        )

    source_ws = _sheet(
        wb,
        "Source Data",
        "Original and normalized values",
        subtitle,
        (
            "Document ID",
            "Statement ID",
            "Field",
            "Original value",
            "Normalized value",
            "Currency",
            "Review status",
            "Source location",
        ),
    )
    for statement in statements:
        for item in statement.values:
            _append(
                source_ws,
                (
                    statement.document_id,
                    statement.id,
                    item.field,
                    item.original_value,
                    item.normalized_value,
                    item.currency or statement.currency,
                    item.status,
                    _refs((item.source_ref,)),
                ),
                numeric=(5,),
            )

    method = _sheet(
        wb,
        "Methodology",
        "How the figures were made",
        subtitle,
        (
            "Measure",
            "Calculation or policy",
            "Missing-data behavior",
        ),
    )
    for formula in FORMULAS.values():
        _append(method, (formula.name, formula.expression, formula.missing_behavior))
    for name, expression in (
        ("free_cash_flow", "Operating cash flow + signed negative capital expenditure"),
        ("cash_conversion_ratio", "Operating cash flow / positive net income"),
        ("net_working_capital", "Current assets - current liabilities"),
        ("dso", "Average receivables / revenue × actual period days"),
        ("dio", "Average inventory / cost of revenue × actual period days"),
        (
            "dpo",
            "Average payables / cost of revenue × actual days; COGS proxies purchases",
        ),
        ("cash_conversion_cycle", "DSO + DIO - DPO"),
    ):
        _append(
            method,
            (
                name,
                expression,
                "Unavailable when required accepted inputs are missing.",
            ),
        )
    _append(
        method,
        (
            "Workbook values",
            "Python calculation snapshots; no hidden Excel formulas or external links.",
            "Unresolved results are blank and labeled unavailable.",
        ),
    )
    _finish(wb)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
